from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_FILE = "Data_Warehouse_Pro_Dashboard_Research_Documentation.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def add_table(document, headers, rows, widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_shading(hdr_cells[i], "D9EAF7")
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in hdr_cells[i].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(10)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[i].paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(10)

    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    return table


def add_caption(document, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)


def add_code_block(document, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F3F6F8")
    set_cell_border(
        cell,
        top={"val": "single", "sz": "6", "color": "9AA6B2"},
        bottom={"val": "single", "sz": "6", "color": "9AA6B2"},
        left={"val": "single", "sz": "6", "color": "9AA6B2"},
        right={"val": "single", "sz": "6", "color": "9AA6B2"},
    )
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    for line_no, line in enumerate(text.splitlines()):
        if line_no > 0:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)
    document.add_paragraph()


def add_diagram_box(document, title, lines):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    add_code_block(document, "\n".join(lines))


def add_bullet(document, text):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def add_numbered(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)


def setup_document():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = document.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Normal"].font.size = Pt(11)

    for style_name, size, color in [
        ("Title", 20, "1F4E79"),
        ("Heading 1", 15, "1F4E79"),
        ("Heading 2", 13, "2F5597"),
        ("Heading 3", 12, "365F91"),
    ]:
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True

    return document


def add_title_page(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(80)
    run = paragraph.add_run("Data Warehouse Pro Dashboard")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 78, 121)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Research-Style Project Documentation")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(28)
    run = paragraph.add_run(
        "A detailed study of ETL processing, dimensional modeling, OLAP operations, "
        "and multidimensional cube analysis using patient admission data"
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(80)
    run = paragraph.add_run("Developed By: Rudransh Pardeshi")
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Project Type: Data Warehousing and Descriptive Data Mining Dashboard")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Prepared as a formal project report")
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    document.add_page_break()


def add_abstract(document):
    document.add_heading("Abstract", level=1)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = paragraph.add_run(
        "This document presents a detailed research-style explanation of the Data Warehouse Pro "
        "Dashboard, a Streamlit-based analytical system built on synthetic patient admission data. "
        "The project demonstrates how raw operational records can be transformed into a structured "
        "data warehouse using ETL processing, dimensional modeling, star schema design, snowflake "
        "normalization, fact constellation modeling, and OLAP analysis. The central analytical "
        "component is a multidimensional cube that aggregates patient satisfaction and length of "
        "stay across service, age group, and quarter. The dashboard enables descriptive data mining "
        "through aggregation, slicing, dicing, and visual pattern discovery."
    )
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

    document.add_heading("Keywords", level=2)
    paragraph = document.add_paragraph()
    run = paragraph.add_run(
        "Data Warehouse, ETL, OLAP, Data Cube, Star Schema, Snowflake Schema, Fact Constellation, "
        "Descriptive Data Mining, Streamlit, Patient Analytics"
    )
    run.italic = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    document.add_page_break()


def add_main_content(document):
    document.add_heading("1. Introduction", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "Data Warehousing is the process of collecting, transforming, organizing, and storing data "
        "for analytical decision-making. Unlike transaction processing systems that focus on day-to-day "
        "record management, a data warehouse focuses on historical analysis, aggregation, and reporting. "
        "This project demonstrates the complete analytical pipeline using patient admission data stored "
        "in patients.csv."
    )

    document.add_heading("1.1 Project Objectives", level=2)
    objectives = [
        "Transform raw patient admission records into analytical warehouse tables.",
        "Design and demonstrate star, snowflake, and fact constellation schemas.",
        "Construct an OLAP cube using service, age group, and quarter dimensions.",
        "Apply slicing and dicing operations for interactive data exploration.",
        "Visualize analytical results through Streamlit and Plotly.",
        "Explain descriptive data mining concepts used in the project.",
    ]
    for item in objectives:
        add_bullet(document, item)

    document.add_heading("2. Technology Stack", level=1)
    add_table(
        document,
        ["Component", "Technology", "Purpose"],
        [
            ["Programming Language", "Python", "Core application logic"],
            ["Dashboard Framework", "Streamlit", "Interactive web interface"],
            ["Data Processing", "Pandas", "Data loading, transformation, and aggregation"],
            ["Numerical Support", "NumPy", "Synthetic billing amount generation"],
            ["Visualization", "Plotly", "Interactive charts and 3D cube visualization"],
            ["Diagram Rendering", "Mermaid", "Flowcharts and schema diagrams in the dashboard"],
            ["Dataset", "patients.csv", "Raw patient admission records"],
        ],
    )

    document.add_heading("3. Dataset Description", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "The source dataset contains 1,000 synthetic patient admission records. Each row represents "
        "one admission event and includes patient identity, demographic information, admission dates, "
        "service category, and satisfaction score."
    )
    add_table(
        document,
        ["Column", "Description"],
        [
            ["patient_id", "Unique identifier for each patient"],
            ["name", "Patient name"],
            ["age", "Patient age"],
            ["arrival_date", "Date of admission or arrival"],
            ["departure_date", "Date of discharge or departure"],
            ["service", "Hospital service used by the patient"],
            ["satisfaction", "Patient satisfaction score"],
        ],
    )

    document.add_heading("4. System Architecture", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "The system follows a layered data warehouse architecture. Raw CSV data enters the ETL layer, "
        "is converted into dimension and fact tables, and is then explored through OLAP operations in "
        "the Streamlit dashboard."
    )
    add_diagram_box(
        document,
        "Figure 1: High-Level Data Warehouse Flow",
        [
            "+----------------+     +----------------+     +----------------------+",
            "|  patients.csv  | --> |  ETL Process   | --> | Data Warehouse Model |",
            "|  Raw Records   |     | Clean/Transform|     | Facts + Dimensions   |",
            "+----------------+     +----------------+     +----------+-----------+",
            "                                                        |",
            "                                                        v",
            "                                      +-------------------------------+",
            "                                      | OLAP Dashboard and Visuals    |",
            "                                      | Cube, Slice, Dice, Charts     |",
            "                                      +-------------------------------+",
        ],
    )

    document.add_heading("5. ETL Methodology", level=1)
    document.add_heading("5.1 Extract", level=2)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "The extraction phase reads patients.csv using Pandas. The raw file behaves like an OLTP-style "
        "operational source where each row stores an individual patient admission record."
    )

    document.add_heading("5.2 Transform", level=2)
    transformations = [
        "Convert arrival_date and departure_date into datetime format.",
        "Create Dim_Date with date_id, year, month, day, and quarter.",
        "Create Dim_Patient with derived age_group values.",
        "Create Dim_Service and map each service to a department.",
        "Create Dim_Department as a normalized higher-level dimension.",
        "Create Fact_Admissions with foreign keys and analytical measures.",
        "Calculate length_of_stay from arrival and departure dates.",
        "Create Fact_Billing with a synthetic total_cost measure.",
    ]
    for item in transformations:
        add_bullet(document, item)

    document.add_heading("5.3 Load", level=2)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "The transformed data is loaded into in-memory warehouse-style tables returned by the "
        "load_and_process_data function. These tables are then consumed by the Streamlit dashboard."
    )
    add_diagram_box(
        document,
        "Figure 2: ETL Block Diagram",
        [
            "Raw Patient Data",
            "       |",
            "       v",
            "Date Parsing --> Feature Engineering --> Key Generation",
            "       |                 |                    |",
            "       v                 v                    v",
            " Dim_Date          Dim_Patient           Dim_Service",
            "                                            |",
            "                                            v",
            "                                      Dim_Department",
            "       |",
            "       v",
            " Fact_Admissions --> Fact_Billing --> Analytical Dashboard",
        ],
    )

    document.add_heading("6. Data Warehouse Design", level=1)
    document.add_heading("6.1 Dimension Tables", level=2)
    add_table(
        document,
        ["Dimension", "Major Columns", "Analytical Purpose"],
        [
            ["Dim_Patient", "patient_id, name, age, age_group", "Analyzes admissions by patient demographics"],
            ["Dim_Date", "date_id, year, month, day, quarter", "Supports time-based aggregation"],
            ["Dim_Service", "service_id, service_name, department_id", "Analyzes data by hospital service"],
            ["Dim_Department", "department_id, department_name", "Represents a normalized service hierarchy"],
        ],
    )

    document.add_heading("6.2 Fact Tables", level=2)
    add_table(
        document,
        ["Fact Table", "Grain", "Measures"],
        [
            ["Fact_Admissions", "One row per patient admission", "length_of_stay, satisfaction"],
            ["Fact_Billing", "One billing record generated from one admission", "total_cost"],
        ],
    )

    document.add_heading("7. Schema Models", level=1)
    document.add_heading("7.1 Star Schema", level=2)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "A star schema places the main fact table at the center and connects it directly to dimension "
        "tables. In this project, Fact_Admissions is connected to Dim_Patient, Dim_Date, and Dim_Service."
    )
    add_diagram_box(
        document,
        "Figure 3: Star Schema",
        [
            "                 +-------------+",
            "                 | Dim_Patient |",
            "                 +------+------+",
            "                        |",
            "+----------+     +------+-------+     +-------------+",
            "| Dim_Date | --> | Fact_Admissions | <-- | Dim_Service |",
            "+----------+     +--------------+     +-------------+",
        ],
    )

    document.add_heading("7.2 Snowflake Schema", level=2)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "A snowflake schema normalizes dimensions into related sub-dimensions. Here, Dim_Service is "
        "linked to Dim_Department, creating a service-to-department hierarchy."
    )
    add_diagram_box(
        document,
        "Figure 4: Snowflake Schema",
        [
            "                 +-------------+",
            "                 | Dim_Patient |",
            "                 +------+------+",
            "                        |",
            "+----------+     +------+-------+     +-------------+     +----------------+",
            "| Dim_Date | --> | Fact_Admissions | <-- | Dim_Service | <-- | Dim_Department |",
            "+----------+     +--------------+     +-------------+     +----------------+",
        ],
    )

    document.add_heading("7.3 Fact Constellation Schema", level=2)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "A fact constellation schema contains multiple fact tables that share common dimensions. "
        "The project demonstrates this by connecting Fact_Admissions and Fact_Billing to the same "
        "patient, date, and service dimensions."
    )
    add_diagram_box(
        document,
        "Figure 5: Fact Constellation Schema",
        [
            "                 +-------------+",
            "                 | Dim_Patient |",
            "                 +------+------+",
            "                        |",
            "       +----------------+----------------+",
            "       |                                 |",
            "+------+-------+                 +-------+------+",
            "| Fact_Admissions |             | Fact_Billing |",
            "+------+-------+                 +-------+------+",
            "       |                                 |",
            "       +---------- Dim_Date -------------+",
            "       +---------- Dim_Service ----------+",
        ],
    )

    document.add_heading("8. OLAP and Data Cube Analysis", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "OLAP, or Online Analytical Processing, is used to analyze data across multiple dimensions. "
        "The central OLAP structure in this project is a three-dimensional cube."
    )
    add_table(
        document,
        ["Cube Component", "Project Field"],
        [
            ["X-axis dimension", "service_name"],
            ["Y-axis dimension", "age_group"],
            ["Z-axis dimension", "quarter"],
            ["Measure 1", "Average satisfaction"],
            ["Measure 2", "Average length_of_stay"],
        ],
    )
    add_diagram_box(
        document,
        "Figure 6: OLAP Cube Construction",
        [
            "Fact_Admissions",
            "      |",
            "      +--> Group by Service",
            "      +--> Group by Age Group",
            "      +--> Group by Quarter",
            "      +--> Aggregate Measures",
            "                  |",
            "                  v",
            "          3D OLAP Data Cube",
            "                  |",
            "                  v",
            "          Plotly 3D Visualization",
        ],
    )

    document.add_heading("8.1 Conceptual Cube Diagram", level=2)
    add_code_block(
        document,
        """                         Quarter
                            ^
                            |
                    Q4  +---------+---------+---------+
                        /         /         /         /|
                   Q3  +---------+---------+---------+ |
                      /         /         /         /| |
                 Q2  +---------+---------+---------+ | |
                    /         /         /         /| | |
               Q1  +---------+---------+---------+ | | |
                   | Surgery |   ICU   |Emergency| | | |
                   +---------+---------+---------+ | | +
                   |  18-34  |  35-54  |   55+   | | /
                   +---------------------------------|/
                              Service ------------->
                  /
                 /
            Age Group"""
    )
    add_caption(document, "Figure 7: Conceptual representation of the service, age group, and quarter cube")

    document.add_heading("8.2 Cube Query Logic", level=2)
    add_code_block(
        document,
        """cube = df_star.groupby(
    ['service_name', 'age_group', 'quarter']
)['satisfaction'].mean().reset_index()"""
    )
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "Each cube cell stores an aggregated measure for one combination of service, age group, and quarter. "
        "For example, one cell may represent the average satisfaction score of ICU patients aged 35-54 in quarter 2."
    )

    document.add_heading("9. Slicing Operation", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "Slicing fixes one dimension of the cube to a single selected value. In this project, the user "
        "selects one service, such as surgery or ICU, and the dashboard displays stay length distribution "
        "for only that service."
    )
    add_diagram_box(
        document,
        "Figure 8: Slice Operation",
        [
            "Full Cube: Service x Age Group x Quarter",
            "                 |",
            "                 v",
            "Select service_name = surgery",
            "                 |",
            "                 v",
            "2D Slice: Age Group x Quarter",
            "                 |",
            "                 v",
            "Length of Stay Histogram",
        ],
    )

    document.add_heading("10. Dicing Operation", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "Dicing filters multiple dimensions at the same time. The dashboard allows users to select multiple "
        "services and multiple age groups, producing a smaller sub-cube for focused comparison."
    )
    add_diagram_box(
        document,
        "Figure 9: Dice Operation",
        [
            "Full Cube: Service x Age Group x Quarter",
            "                 |",
            "                 v",
            "Filter selected services",
            "                 |",
            "                 v",
            "Filter selected age groups",
            "                 |",
            "                 v",
            "Smaller Analytical Sub-Cube",
            "                 |",
            "                 v",
            "Satisfaction vs Stay Length Scatter Plot",
        ],
    )

    document.add_heading("11. Descriptive Data Mining Concepts", level=1)
    add_table(
        document,
        ["Concept", "Use in Project"],
        [
            ["Data Preprocessing", "Dates are parsed and raw fields are prepared for analysis"],
            ["Feature Engineering", "age_group, length_of_stay, quarter, and department are derived"],
            ["Aggregation", "Mean satisfaction and mean stay length are computed"],
            ["Multidimensional Analysis", "Data is analyzed through service, age group, and quarter"],
            ["Filtering", "Slice and dice operations narrow the analytical scope"],
            ["Pattern Discovery", "Charts reveal distribution, relationship, and comparison patterns"],
        ],
    )

    document.add_heading("12. Dashboard Modules", level=1)
    add_numbered(document, "Concepts Overview: displays raw data preview, metrics, and ETL flow.")
    add_numbered(document, "Schema Designs: demonstrates star, snowflake, and fact constellation schemas.")
    add_numbered(document, "OLAP Operations: provides cube, slicing, and dicing interactions.")

    document.add_heading("13. Results and Interpretation", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "The system successfully transforms raw patient admission data into analytical tables and supports "
        "interactive exploration. The cube enables multi-perspective analysis by combining service, age group, "
        "and quarter. Slicing provides single-service analysis, while dicing enables focused comparison between "
        "selected services and age groups. These techniques support descriptive discovery of hospital service "
        "patterns without requiring predictive machine learning."
    )

    document.add_heading("14. Conclusion", level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(
        "The Data Warehouse Pro Dashboard demonstrates a complete data warehouse workflow, from raw source data "
        "to dimensional modeling and OLAP visualization. It provides practical examples of ETL, star schema, "
        "snowflake schema, fact constellation schema, cube aggregation, slicing, dicing, and descriptive data "
        "mining. The project is especially useful for understanding how multidimensional analytical systems are "
        "designed and explored through an interactive dashboard."
    )

    document.add_heading("References", level=1)
    refs = [
        "Kimball, R. and Ross, M. The Data Warehouse Toolkit: The Definitive Guide to Dimensional Modeling.",
        "Inmon, W. H. Building the Data Warehouse.",
        "Pandas Documentation: Data analysis and manipulation library for Python.",
        "Streamlit Documentation: Framework for building interactive data applications.",
        "Plotly Documentation: Interactive graphing library for Python.",
    ]
    for ref in refs:
        add_bullet(document, ref)


def add_headers_footers(document):
    for section in document.sections:
        header = section.header
        paragraph = header.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = paragraph.add_run("Data Warehouse Pro Dashboard")
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)

        footer = section.footer
        paragraph = footer.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run("Research-Style Project Documentation")
        run.font.name = "Times New Roman"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(89, 89, 89)


def main():
    document = setup_document()
    add_title_page(document)
    add_abstract(document)
    document.add_section(WD_SECTION.NEW_PAGE)
    add_main_content(document)
    add_headers_footers(document)
    document.save(OUTPUT_FILE)
    print(OUTPUT_FILE)


if __name__ == "__main__":
    main()
